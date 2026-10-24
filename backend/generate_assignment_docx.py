import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_docx():
    base_dir = r"c:\Users\arshd\OneDrive - vit.ac.in\Desktop\safebite (1)\safebite"
    output_docx_path = os.path.join(base_dir, "Digital_Assignment_3_Review_3_SafeBite.docx")
    copy_target_path = os.path.join(base_dir, "Digital Assignment 3 (1).docx")
    
    doc = Document()
    
    # Configure 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Color Palette
    c_navy = RGBColor(15, 23, 42)      # #0F172A
    c_blue = RGBColor(2, 132, 199)     # #0284C7
    c_body = RGBColor(51, 65, 85)      # #334155
    c_gray = RGBColor(100, 116, 139)   # #64748B
    c_green = RGBColor(5, 150, 105)    # #059669
    
    # Helper to add styled headings
    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(22)
        run.font.color.rgb = c_navy
        run.font.name = "Arial"
        return p

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(16)
        run = p.add_run(text)
        run.font.size = Pt(12)
        run.font.color.rgb = c_blue
        run.font.name = "Arial"
        return p

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(15)
        run.font.color.rgb = c_navy
        run.font.name = "Arial"
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = c_blue
        run.font.name = "Arial"
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = c_navy
        run.font.name = "Arial"
        return p

    def add_p(text, bold_prefix="", italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.size = Pt(9.5)
            r_bold.font.color.rgb = c_navy
            r_bold.font.name = "Arial"
        run = p.add_run(text)
        run.font.size = Pt(9.5)
        run.font.color.rgb = c_body
        run.italic = italic
        run.font.name = "Arial"
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.size = Pt(9.5)
            r_bold.font.color.rgb = c_navy
            r_bold.font.name = "Arial"
        run = p.add_run(text)
        run.font.size = Pt(9.5)
        run.font.color.rgb = c_body
        run.font.name = "Arial"
        return p

    def add_caption(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run(text)
        run.italic = True
        run.font.size = Pt(8.5)
        run.font.color.rgb = c_gray
        run.font.name = "Arial"
        return p

    def insert_image(rel_path, caption_text, width_inches=6.0):
        full_path = os.path.join(base_dir, rel_path)
        if os.path.exists(full_path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run()
            run.add_picture(full_path, width=Inches(width_inches))
            add_caption(f"Figure: {caption_text}")
        else:
            add_p(f"[Image file not found: {rel_path}]", italic=True)

    # ================= COVER / METADATA =================
    add_title("Vellore Institute of Technology, Chennai")
    add_subtitle("School of Computer Science & Engineering | Fall Semester 2026-2027")
    
    # Metadata Header Card (Table)
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="0284C7", sz="8")
    
    meta_rows = [
        ("Course & Code:", "Software Engineering (BCSE301L)"),
        ("Assessment Component:", "Digital Assignment 3 / Review 3 — Project Implementation Evidence"),
        ("Faculty In-Charge:", "Dr. Alok Chauhan"),
        ("Project Title:", "SafeBite — Restaurant Food Safety Transparency Platform"),
        ("Team Members:", "Rudransh Sharma (Reg: 24BAI1510)\nArshdeep Pasricha (Reg: 24BAI1501)"),
        ("GitHub Repository:", "https://github.com/rudransh27sharma/SafeBite")
    ]
    
    for i, (k, v) in enumerate(meta_rows):
        row = table.rows[i]
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=80, bottom=80, left=120, right=120)
        set_cell_margins(c1, top=80, bottom=80, left=120, right=120)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(k)
        r0.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = c_navy
        r0.font.name = "Arial"
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(v)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = c_body
        r1.font.name = "Arial"
        
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ================= EXECUTIVE SUMMARY & PROBLEM STATEMENT =================
    add_h1("1. Project Overview & Problem Statement")
    add_p(
        "Commercial online food delivery aggregators (e.g. Swiggy, Zomato, UberEats) heavily optimize user interfaces "
        "for photographs, discounts, and promotional reviews, but completely obscure critical public health and food safety inspection histories. "
        "Consequently, consumers face severe information asymmetry, unwittingly ordering from facilities with active sanitary violations, "
        "pest infestations, or suspended food licenses. Concurrently, municipal food safety authorities—such as the Food Safety and Standards "
        "Authority of India (FSSAI) and the New York City Department of Health and Mental Hygiene (NYC DOHMH)—maintain official audit registries, "
        "yet these datasets remain siloed, fragmented, and inaccessible to diners during point-of-sale purchasing decisions."
    )
    add_p(
        "SafeBite bridges this gap by deploying an automated intelligence pipeline that ingests verified government inspection records, "
        "maintains strict data provenance, evaluates risk through a deterministic, explainable assessment engine, "
        "and injects real-time safety scores directly into user ordering journeys via a responsive web application, "
        "an interactive companion delivery ordering platform, and a Google Chrome Manifest V3 browser extension."
    )

    # ================= USER STORIES =================
    add_h1("2. User Stories & Persona Requirements")
    add_bullet(" As a health-conscious diner, I want to search restaurants by cuisine and city, viewing safety score rings (0-100) and grade badges (A/B/C/Alert) to make informed dining decisions.", "US-01 (Consumer Discovery):")
    add_bullet(" As a consumer, I want to filter listings exclusively by 'Verified Official Only' (FSSAI/NYC OpenData) so that I can differentiate authentic government audits from demonstration benchmark records.", "US-02 (Provenance Filtering):")
    add_bullet(" As a customer, I want to inspect a restaurant's historical inspection timeline, viewing positive compliance factors, critical infractions, and repeat violation penalties with explicit explanations.", "US-03 (Explainable Assessment):")
    add_bullet(" As a restaurant owner, I want to submit formal Data Correction Requests with supporting reinspection documentation when my listing contains outdated or inaccurate health records.", "US-04 (Owner Dispute):")
    add_bullet(" As an administrator, I want a dedicated moderation dashboard to review, approve, or reject owner correction requests and monitor data freshness and rejected import logs.", "US-05 (Admin Governance):")
    add_bullet(" As an online food delivery customer, I want to see a floating SafeBite Safety Pill injected over the restaurant menu while ordering food, with a pre-checkout advisory modal before payment.", "US-06 (Companion Ordering Demo):")
    add_bullet(" As a Chrome user, I want a browser extension that reads restaurant titles on delivery websites and queries the SafeBite API in real time without manual copy-pasting.", "US-07 (Browser Extension):")

    doc.add_page_break()

    # ================= SECTION 1: TESTING =================
    add_title("Section 1: Integration, Regression, and Mutation Testing")
    add_subtitle("Automated Verification Suite — 23/23 Pytest Passing (100%) | 100% Mutation Score")
    
    add_p(
        "SafeBite enforces a zero-compromise testing strategy across all architectural layers. Food safety scores "
        "directly impact business reputations and consumer health; therefore, all scoring algorithms, data provenance constraints, "
        "matching APIs, and regression bugfixes are systematically verified via automated test suites."
    )
    
    add_h2("1.1 Integration Testing Execution Evidence")
    add_p(
        "Integration tests validate API endpoint contracts, database transactional persistence, fuzzy restaurant matching, "
        "and multi-role administrative review cycles using the FastAPI TestClient and Pytest."
    )
    insert_image(
        "docs/project-evidence/screenshots/12_pytest_test_results.png",
        "Automated Pytest Execution Console — All 23 Integration, Unit, and Regression tests passing (100% Green).",
        width_inches=6.0
    )
    
    add_h3("Comprehensive Test Cases Catalog & Results Table")
    
    test_cases_data = [
        ("TC-AUTH-01", "test_api.py", "User registration, password hashing & JWT token issue", "Valid email/pass", "200 OK, JWT returned", "PASSED"),
        ("TC-DISC-01", "test_api.py", "Restaurant approval & public discovery listing", "Owner submission", "Discoverable in search", "PASSED"),
        ("TC-ENG-01", "test_assessment.py", "Clean inspection evaluates to maximum score (100)", "0 Violations", "Score: 100, Grade: A", "PASSED"),
        ("TC-ENG-02", "test_assessment.py", "Critical violation deduction (-15 pts)", "1 Critical violation", "Score: 85, listed deduction", "PASSED"),
        ("TC-ENG-03", "test_assessment.py", "Stale inspection penalty (>365 days)", "Date 400d ago", "-10 pts & warning pill", "PASSED"),
        ("TC-ENG-04", "test_assessment.py", "Repeat violation 1.5x escalating multiplier", "is_repeat=True", "Penalty: -22.5 pts", "PASSED"),
        ("TC-ENG-05", "test_assessment.py", "Stale inspection warning flag generation", "Inspection > 365d", "Status = 'Stale'", "PASSED"),
        ("TC-ENG-06", "test_assessment.py", "Suspended license override to High Risk Alert", "license_status='Suspended'", "Score <= 35, 'High Risk'", "PASSED"),
        ("TC-ENG-07", "test_assessment.py", "Synthetic record disclaimer generation", "is_synthetic=True", "Disclaimer text emitted", "PASSED"),
        ("TC-ENG-08", "test_assessment.py", "Uninspected restaurant fallback handling", "0 Inspections", "Score: 50, Needs Audit", "PASSED"),
        ("TC-ENG-09", "test_assessment.py", "Verified customer complaint deduction", "1 Verified complaint", "-5 pts deduction", "PASSED"),
        ("TC-CORR-01", "test_disputes_and_admin.py", "Owner dispute submission & admin approval", "Evidence URL + reason", "Status updated 'Approved'", "PASSED"),
        ("TC-MATCH-01", "test_extension_and_matching.py", "Exact name match returns assessment", "Haldiram's, Delhi", "Matched: true, conf >= 0.90", "PASSED"),
        ("TC-MATCH-02", "test_extension_and_matching.py", "Fuzzy restaurant name match", "Chipotle Grill, NYC", "Matched: true, conf >= 0.70", "PASSED"),
        ("TC-MATCH-03", "test_extension_and_matching.py", "POST variant match endpoint", "JSON payload", "200 OK, assessment returned", "PASSED"),
        ("TC-MATCH-04", "test_extension_and_matching.py", "Unmatched restaurant graceful fallback", "NonExistentBakery", "Matched: false, conf 0.0", "PASSED"),
        ("TC-PROV-01", "test_provenance.py", "API returns complete provenance metadata", "GET /api/restaurants/R0001", "source_name, is_synthetic", "PASSED"),
        ("TC-PROV-02", "test_provenance.py", "Synthetic benchmark records explicitly tagged", "1000 Fixtures", "is_synthetic == True", "PASSED"),
        ("TC-PROV-03", "test_provenance.py", "Verified official records marked non-synthetic", "FSSAI / NYC records", "is_synthetic == False", "PASSED"),
        ("TC-REG-01", "test_regression.py", "Admin login succeeds (bcrypt hash fix)", "admin@safebite.demo", "200 OK, JWT returned", "PASSED"),
        ("TC-REG-02", "test_regression.py", "Customer login succeeds with seeded credentials", "customer@safebite.demo", "200 OK, JWT returned", "PASSED"),
        ("TC-REG-03", "test_regression.py", "Database URL resolution not hardcoded to D: drive", "config.DATABASE_URL", "Dynamic path, no 'd:/'", "PASSED"),
        ("TC-REG-04", "test_regression.py", "Profile alias route GET /api/profile returns user", "Bearer JWT Header", "200 OK, User Profile", "PASSED"),
    ]
    
    t_tests = doc.add_table(rows=len(test_cases_data)+1, cols=6)
    t_tests.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_tests, color="CBD5E1", sz="4")
    
    headers = ["Test ID", "Suite", "Purpose / Scope", "Input", "Expected Output", "Status"]
    col_widths = [Inches(0.9), Inches(1.3), Inches(1.8), Inches(1.0), Inches(1.2), Inches(0.8)]
    
    for c_idx, h_text in enumerate(headers):
        cell = t_tests.rows[0].cells[c_idx]
        cell.width = col_widths[c_idx]
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.name = "Arial"
        
    for r_idx, t_row in enumerate(test_cases_data):
        row = t_tests.rows[r_idx+1]
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(t_row):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(8)
            r.font.name = "Arial"
            if c_idx == 5:
                r.bold = True
                r.font.color.rgb = c_green
            elif c_idx == 0:
                r.bold = True
                r.font.color.rgb = c_navy
            else:
                r.font.color.rgb = c_body

    add_h2("1.2 Mutation Testing Evidence (Fault Injection)")
    add_p(
        "Mutation testing evaluates test suite quality by injecting intentional semantic mutations into the scoring logic "
        "(mutating base scores, disabling staleness penalties, altering repeat multipliers, and loosening license revocation caps). "
        "SafeBite's test suite killed 100% of generated mutants (7/7 mutants killed), proving that the tests rigorously catch scoring defects."
    )
    insert_image(
        "docs/project-evidence/screenshots/13_mutation_testing_results.png",
        "Mutation Testing Console Report — 7/7 Mutants Killed (100% Mutation Score).",
        width_inches=6.0
    )

    add_h2("1.3 Regression Testing Analysis")
    add_p(
        "The regression test suite permanently guarantees that baseline defects identified during initial audit do not recur: "
        "(1) REG-001 locks the bcrypt password hash verification fix, (2) REG-002 locks customer credential validation, "
        "(3) REG-003 guarantees dynamic database path resolution on non-D: drive platforms, and (4) REG-004 locks the GET /api/profile session route."
    )

    doc.add_page_break()

    # ================= SECTION 2: VERSION MANAGEMENT & SYSTEM BUILDING =================
    add_title("Section 2: Version Management and System Building")
    add_subtitle("Git Lifecycle, Reproducible Production Builds, and Docker Container Orchestration")

    add_h2("2.1 Version Control Strategy & Git History")
    add_p(
        "SafeBite adheres to a structured Git branching strategy with protected 'main' baseline, dedicated feature branches, "
        "Conventional Commit syntax ('feat:', 'fix:', 'test:', 'refactor:', 'docs:'), and strict .gitignore exclusions protecting secrets."
    )
    insert_image(
        "docs/project-evidence/screenshots/14_git_version_control.png",
        "Git Status & Commit History Console — Clean working tree with structured, semantic commit history.",
        width_inches=6.0
    )

    add_h2("2.2 System Building & Compilation Verification")
    add_p(
        "The frontend production bundle was compiled using Vite 4.5.14 with zero errors, producing optimized chunks: "
        "dist/index.html (0.77 kB), dist/assets/index.css (39.26 kB), and dist/assets/index.js (319 kB). "
        "Backend Python schemas comply fully with Pydantic v2 and SQLAlchemy 2.0 without runtime deprecation warnings."
    )

    add_h2("2.3 Containerized System Startup (Docker Compose)")
    add_p(
        "SafeBite provides a turnkey multi-container setup via docker-compose.yml. The configuration links the FastAPI backend "
        "(exposed on :8000) and the React frontend (exposed on :3000) over an isolated Docker bridge network with automated health checks."
    )
    insert_image(
        "docs/project-evidence/screenshots/15_docker_compose_evidence.png",
        "Docker Compose Orchestration Console — Container health check verification and network bridge startup.",
        width_inches=6.0
    )

    doc.add_page_break()

    # ================= SECTION 3: DEVELOPED FUNCTIONALITIES =================
    add_title("Section 3: Screenshots of Developed Functionalities")
    add_subtitle("Complete User Journey & Interface Verification")

    add_h2("3.1 Landing Page & Discovery Portal")
    add_p(
        "The SafeBite homepage features a hero search bar, direct category exploration, and a transparency banner explaining "
        "how public health inspection ratings are computed. It introduces the user to SafeBite's mission of food safety transparency."
    )
    insert_image(
        "docs/project-evidence/screenshots/01_landing_page.png",
        "SafeBite Landing Page displaying search bar, hero banner, and key feature highlights.",
        width_inches=5.8
    )

    add_h2("3.2 Restaurant Discovery Directory")
    add_p(
        "The restaurant discovery directory (/restaurants) displays search results with real-time safety score rings (0-100), "
        "color-coded letter grade badges (Grade A / Excellent, Grade B / Good, Needs Improvement, High Risk), cuisine filters, and city selection."
    )
    insert_image(
        "docs/project-evidence/screenshots/02_restaurant_discovery.png",
        "Restaurant Discovery View showing score rings, grade badges, and search sidebar.",
        width_inches=5.8
    )

    add_h2("3.3 Data Provenance Filtering (Verified Official vs Synthetic)")
    add_p(
        "To prevent misleading users, SafeBite features a Data Provenance filter. When set to 'Verified Official Only', the system queries only authentic "
        "government inspection records from FSSAI and NYC OpenData, displaying an official green shield badge on each card."
    )
    insert_image(
        "docs/project-evidence/screenshots/03_verified_official_filtering.png",
        "Provenance Filtered View displaying exclusively verified FSSAI and NYC OpenData restaurant records.",
        width_inches=5.8
    )

    add_h2("3.4 Restaurant Detail & Explainable Assessment Engine")
    add_p(
        "The restaurant detail view (/restaurants/R0001) showcases the deterministic assessment engine: starting from a base score of 100, "
        "it presents a transparent breakdown of positive compliance factors, specific violation point deductions, repeat multipliers, and direct links to the official government registry audit."
    )
    insert_image(
        "docs/project-evidence/screenshots/04_restaurant_detail_assessment.png",
        "Restaurant Detail View (Haldiram's) with Explainable Assessment Engine factor breakdown and provenance disclaimer.",
        width_inches=5.8
    )

    add_h2("3.5 Restaurant Owner Data Dispute & Correction Modal")
    add_p(
        "Empowers restaurant owners and consumers to submit formal correction requests if an inspection is misattributed or a reinspection has cleared an infraction. "
        "Submissions require issue categorization, detailed descriptions, and supporting verification URLs."
    )
    insert_image(
        "docs/project-evidence/screenshots/05_dispute_correction_modal.png",
        "Data Dispute / Correction Modal allowing owners to submit official reinspection evidence.",
        width_inches=5.8
    )

    add_h2("3.6 Companion Food Ordering Demonstration (/demo/order)")
    add_p(
        "Because commercial food delivery aggregators (Swiggy, Zomato) lack public injection APIs, SafeBite implements an interactive "
        "companion ordering platform. Evaluators can select restaurants, browse dishes, manage a delivery basket, and view a live floating SafeBite Safety Pill."
    )
    insert_image(
        "docs/project-evidence/screenshots/06_companion_ordering_demo.png",
        "Companion Food Ordering Demonstration (/demo/order) with live basket calculation and floating SafeBite Safety Pill.",
        width_inches=5.8
    )

    add_h2("3.7 Interactive Pre-Checkout Food Safety Advisory Modal")
    add_p(
        "Triggered directly from the ordering demo or Chrome extension, this modal provides an immediate food safety breakdown "
        "before payment, summarizing critical infraction risks, kitchen hygiene compliance, and inspector certification credentials."
    )
    insert_image(
        "docs/project-evidence/screenshots/07_companion_safety_modal.png",
        "Pre-Checkout Food Safety Advisory Modal explaining safety compliance before the diner completes an order.",
        width_inches=5.8
    )

    add_h2("3.8 Role-Based Authentication & Portal Login")
    add_p(
        "SafeBite implements secure role-based JWT authentication with pre-configured quick-fill demo accounts for Administrators, "
        "Customers, and Restaurant Owners, featuring bcrypt password hashing and token expiration."
    )
    insert_image(
        "docs/project-evidence/screenshots/08_login_page.png",
        "Authentication Login Screen with role-based demo accounts and password encryption.",
        width_inches=5.8
    )

    add_h2("3.9 Administrator Moderation Dashboard")
    add_p(
        "The Administrator Dashboard (/admin/dashboard) enables moderators to review new restaurant listing applications, "
        "view inspection data synchronization health, and monitor rejected data import records."
    )
    insert_image(
        "docs/project-evidence/screenshots/09_admin_dashboard.png",
        "Administrator Dashboard showing moderation controls, pending applications, and platform metrics.",
        width_inches=5.8
    )

    add_h2("3.10 Administrator Owner Correction Review Queue")
    add_p(
        "The 4th tab of the Admin Dashboard displays all owner data disputes in real time. Administrators inspect the submitted "
        "evidence URL and either approve the update (recalculating the score) or reject it with moderator audit notes."
    )
    insert_image(
        "docs/project-evidence/screenshots/10_admin_corrections_tab.png",
        "Admin Correction Requests Tab displaying pending owner disputes with Approve and Reject actions.",
        width_inches=5.8
    )

    add_h2("3.11 Mobile Responsive Viewport Verification")
    add_p(
        "SafeBite was tested across mobile device viewports (iPhone 14, 390x844px). The responsive Tailwind layout automatically adapts "
        "navigation, safety score cards, and inspection timelines for seamless handheld operation."
    )
    insert_image(
        "docs/project-evidence/screenshots/11_mobile_responsive_view.png",
        "Mobile Responsive Inspection View (390x844 viewport) confirming responsive handheld accessibility.",
        width_inches=3.6
    )

    add_h2("3.12 Google Chrome Extension (Manifest V3)")
    add_p(
        "The SafeBite Chrome Extension (extension/) operates with a Manifest V3 background service worker and content script. "
        "It scrapes restaurant titles from the active delivery page DOM, queries GET /api/restaurants/match, and injects a floating SafeBite pill "
        "that opens the interactive safety modal directly on the third-party delivery site."
    )

    doc.add_page_break()

    # ================= SECTION 4: TOOLS & TECHNOLOGIES =================
    add_title("Section 4: Tools and Technologies Used")
    add_subtitle("Complete Technology Stack & Operational Inventory")

    tech_inventory = [
        ("Python 3.12.4", "Backend Runtime", "Primary programming language for API, scoring, and data ingestion", "Yes", "Active / Core"),
        ("FastAPI 0.100+", "REST API Framework", "High-performance asynchronous API endpoints with automatic OpenAPI docs", "Yes", "Active / Core"),
        ("Uvicorn 0.40.0", "ASGI Web Server", "Production ASGI server running the FastAPI application on port 8000", "Yes", "Active / Core"),
        ("SQLAlchemy 2.0.51", "Database ORM", "Declarative object-relational mapping, transactions, and relational queries", "Yes", "Active / Core"),
        ("SQLite 3.x", "Embedded Database", "Relational database in database/safebite.db with zero-config portability", "Yes", "Active / Core"),
        ("Passlib (Bcrypt) 1.7.4", "Security & Encryption", "Password hashing with salt generation and timing-attack resistance", "Yes", "Active / Core"),
        ("PyJWT 2.13.0", "Authentication", "Cryptographically signed JSON Web Tokens for role-based session control", "Yes", "Active / Core"),
        ("RapidFuzz 3.14.5", "String Matching", "Levenshtein distance fuzzy matching for restaurant name queries", "Yes", "Active / Core"),
        ("React 18.2.0", "Frontend Framework", "Component-based single-page reactive user interface", "Yes", "Active / Core"),
        ("Vite 4.5.14", "Build Tool / Bundler", "Lightning-fast HMR dev server and optimized production bundler", "Yes", "Active / Core"),
        ("Tailwind CSS 3.3.3", "Styling Engine", "Utility-first responsive CSS styling with custom SafeBite design tokens", "Yes", "Active / Core"),
        ("Lucide React 0.268.0", "Iconography", "Consistent SVG icon library for UI controls and food safety shields", "Yes", "Active / Core"),
        ("Axios 1.5.0", "HTTP Client", "Browser HTTP client with automatic JWT bearer authorization interceptors", "Yes", "Active / Core"),
        ("Manifest V3", "Browser Extension", "Chrome extension architecture with service worker and content script", "Yes", "Active / Core"),
        ("Pytest 9.0.3", "Test Automation", "Automated unit, integration, and regression testing framework (23 tests)", "Yes", "Active / Core"),
        ("FastAPI TestClient", "Integration Testing", "In-memory Starlette HTTP client testing API routes without socket binding", "Yes", "Active / Core"),
        ("Custom AST Runner", "Mutation Testing", "Systematic fault injection evaluating scoring engine test resilience", "Yes", "Active / Core"),
        ("Selenium WebDriver 4.45", "Browser Automation", "Automated headless Chrome screenshot and UI validation pipeline", "Yes", "Active / Core"),
        ("ReportLab 5.0.1", "Dossier Generation", "Programmatic PDF compilation generating the 11-page evidence dossier", "Yes", "Active / Core"),
        ("Docker Compose 3.8", "Containerization", "Multi-container orchestration for FastAPI and React Vite SPA", "Yes", "Active / Core"),
        ("Git 2.x", "Version Control", "Distributed version control with semantic branching and commit conventions", "Yes", "Active / Core"),
    ]

    t_tech = doc.add_table(rows=len(tech_inventory)+1, cols=5)
    t_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_tech, color="CBD5E1", sz="4")
    
    t_headers = ["Technology", "Layer", "Purpose in SafeBite", "Local Run?", "Status"]
    t_widths = [Inches(1.4), Inches(1.2), Inches(2.3), Inches(0.8), Inches(0.8)]
    
    for c_idx, h_text in enumerate(t_headers):
        cell = t_tech.rows[0].cells[c_idx]
        cell.width = t_widths[c_idx]
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.name = "Arial"
        
    for r_idx, t_row in enumerate(tech_inventory):
        row = t_tech.rows[r_idx+1]
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(t_row):
            cell = row.cells[c_idx]
            cell.width = t_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(8)
            r.font.name = "Arial"
            if c_idx == 0:
                r.bold = True
                r.font.color.rgb = c_navy
            elif c_idx == 4:
                r.bold = True
                r.font.color.rgb = c_green
            else:
                r.font.color.rgb = c_body

    add_h2("Deliberate External Integration Boundaries")
    add_bullet("Commercial delivery platforms (Swiggy, Zomato, DoorDash) do not expose public third-party APIs for safety overlay injection without enterprise partnership agreements. SafeBite provides a 100% functional, realistic simulation via the Manifest V3 Browser Extension and the Companion Ordering Demo at /demo/order.", "Commercial Delivery APIs:")
    add_bullet("Government portals (FSSAI FoSCoS) do not offer real-time push webhooks for public developers. SafeBite ingests verified official government records via structured JSON batch imports with provenance tracking.", "Government Health Registries:")

    doc.add_page_break()

    # ================= SECTION 5: APPENDIX & ARCHITECTURE DIAGRAMS =================
    add_title("Section 5: System Architecture & Design Diagrams (Appendix)")
    add_subtitle("Software Engineering UML, ER, and Architecture Models")

    add_h2("5.1 System Architecture Diagram")
    add_p(
        "The SafeBite system architecture illustrates the end-to-end data flow: Data ingestion from FSSAI and NYC OpenData into SQLite, "
        "processed through the SQLAlchemy ORM and the Explainable Assessment Engine, exposed via FastAPI REST endpoints, and consumed "
        "by the React Vite SPA, Companion Ordering Demo, and Chrome Browser Extension."
    )
    insert_image(
        "design/architecture/SafeBite_Architecture.drawio.png",
        "SafeBite End-to-End System Architecture Diagram.",
        width_inches=6.0
    )

    add_h2("5.2 Component Diagram")
    add_p(
        "The component diagram details internal subsystems: Auth Module, Restaurant Service, Assessment Engine, "
        "Inspection Service, Ingestion Pipeline, and the Browser Extension Content Script bridge."
    )
    insert_image(
        "design/architecture/component.drawio.png",
        "SafeBite Component Architecture Diagram.",
        width_inches=6.0
    )

    add_h2("5.3 Entity-Relationship (ER) Diagram")
    add_p(
        "The relational data model comprises Users (with roles), Restaurants (with provenance attributes), Inspections, "
        "Violations, Inspectors, Complaints, and Correction Requests (for owner disputes)."
    )
    insert_image(
        "design/architecture/ER.drawio.png",
        "SafeBite Relational Entity-Relationship (ER) Diagram.",
        width_inches=6.0
    )

    add_h2("5.4 Class Diagram")
    add_p(
        "The structural class diagram defines the object-oriented abstractions, schemas, Pydantic data models, and ORM entity relationships."
    )
    insert_image(
        "design/architecture/Class Diagram.drawio.png",
        "SafeBite Structural Class Diagram.",
        width_inches=6.0
    )

    add_h2("5.5 Use Case Diagram")
    add_p(
        "Depicts system actors (Customer, Restaurant Owner, Health Inspector, Administrator) and their interactions with core platform features."
    )
    insert_image(
        "design/architecture/usecase.drawio.png",
        "SafeBite System Use Case Diagram.",
        width_inches=6.0
    )

    add_h2("5.6 Sequence Diagram")
    add_p(
        "Traces the chronological message sequence during restaurant matching, assessment evaluation, and safety modal rendering."
    )
    insert_image(
        "design/architecture/sequence.drawio.png",
        "SafeBite Restaurant Assessment Sequence Diagram.",
        width_inches=6.0
    )

    doc.add_page_break()

    # ================= SECTION 6: INSTRUCTIONS FOR REPRODUCING DEMO =================
    add_title("Section 6: Instructions for Reproducing the Project Demo")
    add_subtitle("Step-by-Step Evaluator Reproduction Walkthrough")

    add_p(
        "Follow these exact steps to reproduce and verify the SafeBite platform from source code:"
    )

    steps = [
        ("Step 1: Install Dependencies", "cd backend && pip install -r requirements.txt reportlab selenium\ncd ../frontend && npm install"),
        ("Step 2: Database Migration & Data Ingestion", "cd backend\npython migrate_db.py\npython ingest_inspection_data.py"),
        ("Step 3: Run Automated Test Suites", "cd backend\npytest -v                          # 23 Automated tests pass (100% Green)\npython tests/run_mutation_testing.py # 7/7 Mutants killed (100% Score)"),
        ("Step 4: Start Backend & Frontend", "Terminal 1 (Backend):  python -m uvicorn app.main:app --host 127.0.0.1 --port 8000\nTerminal 2 (Frontend): cd frontend && npm run dev (or npx vite preview)"),
        ("Step 5: Open Application & Demo", "1. Open Browser: http://localhost:5173/\n2. Filter by Provenance: http://localhost:5173/restaurants (Select 'Verified Official')\n3. Restaurant Detail: http://localhost:5173/restaurants/R0001 (Explainable score breakdown)\n4. Submit Owner Dispute: Click 'Submit Data Correction / Dispute'\n5. Companion Ordering Demo: http://localhost:5173/demo/order (Add items, click 'View Safety Audit Breakdown')\n6. Admin Moderation: http://localhost:5173/admin/dashboard (admin@safebite.demo / adminpassword)")
    ]

    for title, cmd in steps:
        add_h3(title)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(cmd)
        r.font.name = "Courier New"
        r.font.size = Pt(8.5)
        r.font.color.rgb = c_navy

    # Save document to primary target
    doc.save(output_docx_path)
    print(f"Generated successfully: {output_docx_path}")
    print(f"File size: {os.path.getsize(output_docx_path)} bytes")
    
    # Also overwrite 'Digital Assignment 3 (1).docx' as requested
    doc.save(copy_target_path)
    print(f"Updated successfully: {copy_target_path}")
    print(f"File size: {os.path.getsize(copy_target_path)} bytes")

if __name__ == "__main__":
    build_docx()
